import os
import io
import pytest
import fitz
import docx
from fastapi.testclient import TestClient
from app.modules.ai_engine.mock_provider import MockProvider
from app.main import app
from app.services.conversion_storage import conversion_storage

client = TestClient(app)


@pytest.fixture(autouse=True)
def override_ai_provider(monkeypatch):
    import app.api.v1.endpoints.process as process_endpoint
    import app.services.content_processor as content_processor_module

    monkeypatch.setattr(process_endpoint, "get_ai_provider", lambda: MockProvider())
    monkeypatch.setattr(content_processor_module, "get_ai_provider", lambda: MockProvider())


def test_pdf_conversion_and_download():
    """
    Test converting a PDF to Markdown (.md) artifact and downloading it via GET /conversions/{id}/download.
    """
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "PDF Technical Specification Document Header")
    page.insert_text((50, 80), "This PDF content should be normalized into Markdown.")

    buffer = io.BytesIO()
    doc.save(buffer)
    pdf_bytes = buffer.getvalue()
    doc.close()

    response = client.post(
        "/api/v1/convert",
        files=[("files", ("tech-spec.pdf", pdf_bytes, "application/pdf"))],
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert len(data["conversions"]) == 1

    item = data["conversions"][0]
    assert item["original_filename"] == "tech-spec.pdf"
    assert item["markdown_filename"] == "tech-spec.md"
    assert item["output_format"] == "markdown"
    assert item["id"] != ""

    # Test downloading the generated Markdown artifact
    download_res = client.get(f"/api/v1/conversions/{item['id']}/download")
    assert download_res.status_code == 200
    assert "text/markdown" in download_res.headers["content-type"]
    assert 'attachment; filename="tech-spec.md"' in download_res.headers["content-disposition"]
    md_content = download_res.text
    assert "PDF Technical Specification Document Header" in md_content


def test_docx_conversion_and_download():
    """
    Test converting a DOCX file to Markdown (.md) artifact and downloading it.
    """
    doc = docx.Document()
    doc.add_heading("Financial Report 2026", level=1)
    doc.add_paragraph("Executive summary of financial performance.")

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    response = client.post(
        "/api/v1/convert",
        files=[("files", ("financial.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    item = data["conversions"][0]
    assert item["markdown_filename"] == "financial.md"

    download_res = client.get(f"/api/v1/conversions/{item['id']}/download")
    assert download_res.status_code == 200
    assert "Financial Report 2026" in download_res.text


def test_txt_conversion_and_download():
    """
    Test converting plain text file to Markdown.
    """
    sample_text = "Title: Project Plan\n\n- Task 1: Research\n- Task 2: Build\n"
    response = client.post(
        "/api/v1/convert",
        files=[("files", ("plan.txt", sample_text.encode("utf-8"), "text/plain"))],
    )
    assert response.status_code == 200
    data = response.json()
    item = data["conversions"][0]
    assert item["markdown_filename"] == "plan.md"

    download_res = client.get(f"/api/v1/conversions/{item['id']}/download")
    assert download_res.status_code == 200
    assert "Project Plan" in download_res.text


def test_multiple_files_conversion():
    """
    Verify uploading multiple files produces separate individual .md artifacts and a combined artifact.
    """
    file_a = ("doc1.txt", b"Content of document A", "text/plain")
    file_b = ("doc2.txt", b"Content of document B", "text/plain")

    response = client.post(
        "/api/v1/convert",
        files=[("files", file_a), ("files", file_b)],
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert len(data["conversions"]) == 2
    assert data["combined_conversion_id"] is not None

    # Test downloading combined artifact
    combined_id = data["combined_conversion_id"]
    combined_res = client.get(f"/api/v1/conversions/{combined_id}/download")
    assert combined_res.status_code == 200
    assert "Source: doc1.txt" in combined_res.text
    assert "Source: doc2.txt" in combined_res.text


def test_large_file_conversion():
    """
    Verify conversion of a 20 MB text document works cleanly within configured limits.
    """
    # Create ~20 MB payload (20 * 1024 * 1024 bytes)
    chunk = "Large document block line with research information.\n" * 100
    repeats = (20 * 1024 * 1024) // len(chunk)
    large_text = chunk * repeats
    large_bytes = large_text.encode("utf-8")

    response = client.post(
        "/api/v1/convert",
        files=[("files", ("large-research.txt", large_bytes, "text/plain"))],
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    item = data["conversions"][0]
    assert item["source_size_bytes"] >= 20000000
    assert item["markdown_filename"] == "large-research.md"

    # Download large artifact
    download_res = client.get(f"/api/v1/conversions/{item['id']}/download")
    assert download_res.status_code == 200
    assert len(download_res.text) >= 20000000


def test_fallback_generated_markdown_download(monkeypatch):
    """
    Verify that fallback-extracted Markdown (e.g. PyMuPDF or OCR) is also downloadable.
    """
    from app.modules.ingestion.markitdown_service import markitdown_service
    from app.schemas.processing import ExtractionResult, ExtractionStatus

    async def mock_failed_convert(file_bytes, filename):
        return ExtractionResult(
            filename=filename,
            file_type="pdf",
            text="",
            metadata={},
            extraction_status=ExtractionStatus.FAILED,
            error_message="Simulated conversion failure",
            extraction_method="markitdown",
            output_format="markdown",
            fallback_used=False,
        )

    monkeypatch.setattr(markitdown_service, "convert", mock_failed_convert)

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Fallback PDF Content")
    buffer = io.BytesIO()
    doc.save(buffer)
    pdf_bytes = buffer.getvalue()
    doc.close()

    response = client.post(
        "/api/v1/convert",
        files=[("files", ("fallback_test.pdf", pdf_bytes, "application/pdf"))],
    )
    assert response.status_code == 200
    data = response.json()
    item = data["conversions"][0]
    assert item["fallback_used"] is True

    download_res = client.get(f"/api/v1/conversions/{item['id']}/download")
    assert download_res.status_code == 200
    assert "Fallback PDF Content" in download_res.text


def test_download_security_path_traversal():
    """
    Verify path traversal attempts are strictly rejected.
    """
    invalid_ids = [
        "../etc/passwd",
        "..\\..\\windows\\system32",
        "8e4b/../../secret",
        "invalid_id_not_uuid",
    ]

    for inv_id in invalid_ids:
        res = client.get(f"/api/v1/conversions/{inv_id}/download")
        assert res.status_code in [400, 404]


def test_process_with_conversion_ids():
    """
    Verify AI processing using pre-converted conversion_ids.
    """
    # 1. Convert file first
    conv_res = client.post(
        "/api/v1/convert",
        files=[("files", ("specs.txt", b"System launch date is October 15, 2026.", "text/plain"))],
    )
    assert conv_res.status_code == 200
    conv_id = conv_res.json()["conversions"][0]["id"]

    # 2. Process using conversion_id
    proc_res = client.post(
        "/api/v1/process",
        data={
          "mode": "strict",
          "prompt": "What is the system launch date?",
          "conversion_ids": [conv_id],
        },
    )
    assert proc_res.status_code == 200
    proc_data = proc_res.json()
    assert proc_data["success"] is True
    assert proc_data["mode"] == "strict"
    assert "System launch date is October 15, 2026" in proc_data["source_markdown"]
