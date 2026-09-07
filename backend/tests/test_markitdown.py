import pytest
import io
import fitz
import docx
from PIL import Image

from app.modules.ingestion.service import ingestion_service
from app.modules.ingestion.markitdown_service import markitdown_service
from app.schemas.processing import ExtractionStatus
from app.modules.prompts.builder import prompt_builder


@pytest.mark.asyncio
async def test_txt_to_markdown():
    sample_text = "Project Title\n\nThis is a simple plain text file for testing."
    bytes_data = sample_text.encode("utf-8")

    res = await ingestion_service.process_file("test_doc.txt", bytes_data)
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert res.output_format == "markdown"
    assert "Project Title" in res.text
    assert res.file_type == "txt"


@pytest.mark.asyncio
async def test_md_to_markdown():
    sample_md = "# Architecture\n\n- Component 1\n- Component 2\n"
    bytes_data = sample_md.encode("utf-8")

    res = await ingestion_service.process_file("architecture.md", bytes_data)
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "# Architecture" in res.text
    assert res.output_format == "markdown"


@pytest.mark.asyncio
async def test_docx_to_markdown():
    doc = docx.Document()
    doc.add_heading("Release Notes v1.0", level=1)
    doc.add_paragraph("Features included in this release:")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Feature"
    table.rows[0].cells[1].text = "Status"
    table.rows[1].cells[0].text = "MarkItDown Ingestion"
    table.rows[1].cells[1].text = "Complete"

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    res = await ingestion_service.process_file("release.docx", docx_bytes)
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "Release Notes v1.0" in res.text
    assert res.output_format == "markdown"


@pytest.mark.asyncio
async def test_pdf_to_markdown():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "PDF Document Ingestion Test Header")
    page.insert_text((50, 80), "This PDF will be normalized to Markdown.")

    buffer = io.BytesIO()
    doc.save(buffer)
    pdf_bytes = buffer.getvalue()
    doc.close()

    res = await ingestion_service.process_file("spec.pdf", pdf_bytes)
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "PDF Document Ingestion Test Header" in res.text
    assert res.output_format == "markdown"


@pytest.mark.asyncio
async def test_markitdown_fallback_trigger(monkeypatch):
    """
    Test explicit fallback execution path when MarkItDown fails or returns empty text.
    """
    from app.schemas.processing import ExtractionResult, ExtractionStatus

    async def mock_failed_convert(file_bytes, filename):
        return ExtractionResult(
            filename=filename,
            file_type="pdf",
            text="",
            metadata={},
            extraction_status=ExtractionStatus.FAILED,
            error_message="Simulated conversion failure",
            extraction_method="markitdown",
            output_format="markdown",
            fallback_used=False,
        )

    monkeypatch.setattr(markitdown_service, "convert", mock_failed_convert)

    # Create valid PDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Fallback PDF Content")
    buffer = io.BytesIO()
    doc.save(buffer)
    pdf_bytes = buffer.getvalue()
    doc.close()

    res = await ingestion_service.process_file("fallback_test.pdf", pdf_bytes)
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert res.fallback_used is True
    assert res.extraction_method in ["pymupdf", "ocr"]
    assert "Fallback PDF Content" in res.text


@pytest.mark.asyncio
async def test_image_ocr_path():
    img = Image.new("RGB", (200, 100), color="white")
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    img_bytes = buffer.getvalue()

    res = await ingestion_service.process_file("diagram.png", img_bytes)
    assert res.filename == "diagram.png"
    assert res.file_type == "png"
    assert res.output_format == "markdown"


@pytest.mark.asyncio
async def test_multi_file_combining():
    file_a = ("doc1.txt", "Content of File A".encode("utf-8"))
    file_b = ("doc2.md", "# Title B\nContent of File B".encode("utf-8"))

    combined_text, results, combined_id = await ingestion_service.process_files_and_combine([file_a, file_b])

    assert len(results) == 2
    assert combined_id is not None
    assert "# Source: doc1.txt" in combined_text
    assert "# Source: doc2.md" in combined_text
    assert "---" in combined_text
    assert "Content of File A" in combined_text
    assert "Content of File B" in combined_text


@pytest.mark.asyncio
async def test_empty_file_validation():
    res = await ingestion_service.process_file("empty.txt", b"")
    assert res.extraction_status in [ExtractionStatus.FAILED, ExtractionStatus.PARTIAL]
    assert res.text == ""


def test_prompt_builder_decoupling():
    """
    Verify Prompt Builder receives normalized Markdown without knowing extraction details.
    """
    markdown_content = "# Source: report.pdf\n\n## Section 1\nNormalized Markdown text."
    system_inst, formatted_prompt = prompt_builder.build_prompt(
        mode="strict",
        user_prompt="Summarize Section 1",
        source_content=markdown_content,
    )

    assert "STRICT MODE" in system_inst
    assert "Source Content:\n# Source: report.pdf" in formatted_prompt
