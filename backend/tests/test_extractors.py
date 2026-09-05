import pytest
import io
import fitz
import docx
from PIL import Image
from app.modules.ingestion.extractors.text_extractor import TextExtractor
from app.modules.ingestion.extractors.docx_extractor import DocxExtractor
from app.modules.ingestion.extractors.pdf_extractor import PdfExtractor
from app.modules.ingestion.extractors.image_extractor import ImageExtractor
from app.modules.ingestion.ocr import ocr_engine
from app.modules.ingestion.service import ingestion_service
from app.schemas.processing import ExtractionStatus


def test_text_extraction():
    extractor = TextExtractor()
    sample_content = "Hello, world!\nThis is plain text."
    bytes_data = sample_content.encode("utf-8")

    res = extractor.extract(bytes_data, "test.txt")
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "Hello, world!" in res.text
    assert res.file_type == "txt"
    assert res.metadata["word_count"] == 6


def test_markdown_extraction():
    extractor = TextExtractor()
    sample_content = "# Title\n- Item 1\n- Item 2"
    bytes_data = sample_content.encode("utf-8")

    res = extractor.extract(bytes_data, "notes.md")
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "# Title" in res.text
    assert res.file_type == "md"


def test_docx_extraction():
    extractor = DocxExtractor()
    
    # Create sample docx in memory
    doc = docx.Document()
    doc.add_paragraph("Sample Paragraph 1")
    doc.add_paragraph("Sample Paragraph 2")
    
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Header A"
    hdr_cells[1].text = "Header B"

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    res = extractor.extract(docx_bytes, "sample.docx")
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "Sample Paragraph 1" in res.text
    assert "Header A | Header B" in res.text
    assert res.file_type == "docx"


def test_pdf_text_extraction():
    extractor = PdfExtractor()

    # Create sample PDF with PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "PDF Text Content Test Page 1")

    buffer = io.BytesIO()
    doc.save(buffer)
    pdf_bytes = buffer.getvalue()
    doc.close()

    res = extractor.extract(pdf_bytes, "test.pdf")
    assert res.extraction_status == ExtractionStatus.SUCCESS
    assert "PDF Text Content Test Page 1" in res.text
    assert res.metadata["total_pages"] == 1


def test_image_ocr_abstraction():
    extractor = ImageExtractor()

    # Create sample RGB image in memory
    img = Image.new("RGB", (200, 100), color="white")
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    img_bytes = buffer.getvalue()

    res = extractor.extract(img_bytes, "sample.png")
    assert res.filename == "sample.png"
    assert res.file_type == "png"
    assert res.metadata["width"] == 200


def test_invalid_file_handling():
    is_valid, err = ingestion_service.validate_file("file.exe", b"binary content")
    assert is_valid is False
    assert "Unsupported file type" in err

    res = ingestion_service.process_file("invalid.xyz", b"abc")
    assert res.extraction_status == ExtractionStatus.FAILED
    assert "Unsupported file type" in res.error_message
