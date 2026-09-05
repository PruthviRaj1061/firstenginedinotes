import io
import os
import docx
from app.modules.ingestion.extractors.base import BaseExtractor
from app.schemas.processing import ExtractionResult, ExtractionStatus


class DocxExtractor(BaseExtractor):
    """
    Extractor for Microsoft Word (.docx) documents using python-docx.
    """

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        try:
            doc_file = io.BytesIO(file_bytes)
            doc = docx.Document(doc_file)

            extracted_lines = []

            # 1. Paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    extracted_lines.append(p_text)

            # 2. Tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        extracted_lines.append(" | ".join(row_cells))

            full_text = "\n".join(extracted_lines).strip()
            word_count = len(full_text.split()) if full_text else 0

            return ExtractionResult(
                filename=filename,
                file_type="docx",
                text=full_text,
                metadata={
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables),
                    "char_count": len(full_text),
                    "word_count": word_count,
                },
                extraction_status=ExtractionStatus.SUCCESS if full_text else ExtractionStatus.PARTIAL,
            )
        except Exception as e:
            return ExtractionResult(
                filename=filename,
                file_type="docx",
                text="",
                metadata={},
                extraction_status=ExtractionStatus.FAILED,
                error_message=f"Failed to parse DOCX file: {str(e)}",
            )
